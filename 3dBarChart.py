from openpyxl import Workbook
from openpyxl.chart import BarChart3D, Series, Reference
from openpyxl import load_workbook
wb = load_workbook(filename="test.xlsx")
sheet = wb.active

chart = BarChart3D()
values = Reference(worksheet=sheet,
                   min_row=1,
                   max_row=15,
                   min_col=2,
                   max_col=13)

chart.add_data(values, titles_from_data=True)
sheet.add_chart(chart, "E3")
wb.save("MyChart2.xlsx")